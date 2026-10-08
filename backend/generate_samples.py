import os
import docx
from pypdf import PdfWriter

os.makedirs("sample_documents", exist_ok=True)

# 1. Create Sample DOCX
doc = docx.Document()
doc.add_heading("IN THE HIGH COURT OF JUSTICE", level=1)
doc.add_heading("COMMERCIAL DIVISION - ACCRA", level=2)
doc.add_paragraph("SUIT NO: HC/2026/ACCRA/094")
doc.add_paragraph("BETWEEN:")
doc.add_paragraph("[PLAINTIFF NAME]                          ... PLAINTIFF")
doc.add_paragraph("AND")
doc.add_paragraph("KOFI MENSAH                                ... DEFENDANT")
doc.add_heading("CONSENT ORDER AND PAYMENT TERMS", level=2)
doc.add_paragraph(
    "UPON HEARING Counsel for the Plaintiff and the Defendant agreeing to terms of settlement, "
    "IT IS HEREBY ORDERED BY CONSENT that:"
)
doc.add_paragraph("1. The Defendant shall pay to the Plaintiff the total sum of USD 24,000.00 in full settlement.")
doc.add_paragraph("2. The amount shall be liquidated in 6 equal consecutive monthly instalments of USD 4,000.00 each.")
doc.add_paragraph("3. The first instalment shall become due and payable on the 1st day of next month.")
doc.add_paragraph("4. An agreed interest rate of 5.0% per annum shall apply on the outstanding principal balance.")
doc.add_paragraph("5. In the event of default on any instalment exceeding a grace period of 7 days, a default penalty of 2.0% shall immediately apply on all overdue arrears.")
doc.add_paragraph("6. All payments shall be made in USD or equivalent currency as agreed.")
doc.save("sample_documents/sample_court_order.docx")
print("Created sample_documents/sample_court_order.docx")

# 2. Create Sample Text Document
sample_text = """IN THE HIGH COURT OF JUSTICE
SUIT NO: HC/2026/ACCRA/102
BETWEEN:
[PLAINTIFF NAME] (Plaintiff)
AND
Ama Serwaa (Defendant)

TERMS OF SETTLEMENT AND CONSENT JUDGMENT
1. The Defendant admits liability in the sum of USD 18,000.00.
2. The Defendant shall liquidate the judgment debt in 3 equal monthly instalments of USD 6,000.00.
3. The first instalment is due on 2026-02-01.
4. Interest of 8.0% per annum shall accrue on any unpaid balance.
5. In the event of default past 5 days grace period, a 5.0% penalty shall apply on the arrears.
"""
with open("sample_documents/sample_settlement_agreement.txt", "w", encoding="utf-8") as f:
    f.write(sample_text)
print("Created sample_documents/sample_settlement_agreement.txt")
