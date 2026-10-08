import os
import io
import docx
from pypdf import PdfReader
import pdfplumber
from PIL import Image

def extract_text_from_pdf(file_path: str) -> str:
    text = ""
    # Try pdfplumber first for high-fidelity text extraction
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except Exception:
        pass

    # Fallback to pypdf if pdfplumber didn't extract enough
    if not text.strip():
        try:
            reader = PdfReader(file_path)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        except Exception as e:
            text = f"Error reading PDF: {str(e)}"
            
    return text.strip()

def extract_text_from_docx(file_path: str) -> str:
    try:
        doc = docx.Document(file_path)
        full_text = []
        for para in doc.paragraphs:
            if para.text:
                full_text.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text.append(" | ".join(row_text))
        return "\n".join(full_text)
    except Exception as e:
        return f"Error reading DOCX: {str(e)}"

def extract_text_from_image(file_path: str) -> str:
    try:
        with Image.open(file_path) as img:
            width, height = img.size
            format_name = img.format
            return f"[Image Document: {os.path.basename(file_path)}, Format: {format_name}, Dimensions: {width}x{height}]"
    except Exception as e:
        return f"Error inspecting image: {str(e)}"

def extract_document_text(file_path: str, content_type: str = "") -> str:
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext == ".pdf" or "pdf" in content_type:
        return extract_text_from_pdf(file_path)
    elif ext in [".docx", ".doc"] or "word" in content_type:
        return extract_text_from_docx(file_path)
    elif ext in [".png", ".jpg", ".jpeg"] or "image" in content_type:
        return extract_text_from_image(file_path)
    else:
        # Fallback raw read
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        except Exception as e:
            return f"Unsupported file type: {ext}"
